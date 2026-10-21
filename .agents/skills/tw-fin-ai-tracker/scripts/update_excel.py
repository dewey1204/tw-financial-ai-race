import sys
import json
from copy import copy
from datetime import datetime, timezone, timedelta
try:
    from zoneinfo import ZoneInfo
except ImportError:
    ZoneInfo = None
import openpyxl

# 表列 14 家金控及名稱關鍵字對照（依「分類說明」工作表指定優先順序排列）
HOLDING_KEYWORD_RULES = [
    ("元大金", ["元大"]),
    ("富邦金", ["富邦"]),
    ("國泰金", ["國泰", "國泰世華"]),
    ("中信金", ["中信", "中國信託", "台灣人壽"]),
    ("凱基金", ["凱基", "開發"]),
    ("永豐金", ["永豐"]),
    ("台新新光金", ["台新", "新光"]),
    ("玉山金", ["玉山"]),
    ("兆豐金", ["兆豐"]),
    ("第一金", ["第一"]),
    ("華南金", ["華南"]),
    ("合庫金", ["合庫"]),
    ("國票金", ["國票"]),
    ("臺灣金", ["臺銀", "臺灣銀行", "臺灣金控"]),
]

VALID_HOLDINGS = {holding for holding, _ in HOLDING_KEYWORD_RULES}


def get_taipei_today():
    """取得台灣台北時區 (Asia/Taipei, UTC+8) 的當天日期 (datetime 物件，時間為 00:00:00)"""
    if ZoneInfo is not None:
        try:
            now_tpe = datetime.now(ZoneInfo("Asia/Taipei"))
            return datetime(now_tpe.year, now_tpe.month, now_tpe.day)
        except Exception:
            pass
    tpe_tz = timezone(timedelta(hours=8))
    now_tpe = datetime.now(tpe_tz)
    return datetime(now_tpe.year, now_tpe.month, now_tpe.day)


def parse_update_date(date_val, default_dt):
    """解析傳入的日期字串 (YYYY-MM-DD 或 YYYY/MM/DD)，若未傳入則使用台北時區當天日期"""
    if not date_val:
        return default_dt
    if isinstance(date_val, datetime):
        return datetime(date_val.year, date_val.month, date_val.day)
    text = str(date_val).strip().replace("/", "-")
    for fmt in ("%Y-%m-%d", "%Y%m%d"):
        try:
            parsed = datetime.strptime(text, fmt)
            return datetime(parsed.year, parsed.month, parsed.day)
        except ValueError:
            continue
    return default_dt


def resolve_holding(company_name, raw_holding=""):
    """
    判定機構所屬金控：
    1. 若已傳入表列 14 家金控名稱之一，直接採用。
    2. 若傳入完整金控名稱（如「富邦金控」），或未傳入／非表列名稱，依表列 14 組關鍵字順序比對機構名稱與傳入值。
    3. 若均非屬表列 14 家金控，一律回傳 "未歸類"。
    """
    company_str = str(company_name or "").strip()
    holding_str = str(raw_holding or "").strip()

    if holding_str in VALID_HOLDINGS:
        return holding_str

    if holding_str in ("未歸類", "未分類"):
        # 仍檢查機構名稱是否明確含有表列金控關鍵字，若無則回傳 "未歸類"
        for holding, keywords in HOLDING_KEYWORD_RULES:
            if any(kw in company_str for kw in keywords):
                return holding
        return "未歸類"

    # 依「分類說明」表列順序比對機構名稱與傳入之金控文字
    search_target = f"{company_str} {holding_str}"
    for holding, keywords in HOLDING_KEYWORD_RULES:
        if any(kw in search_target for kw in keywords):
            return holding

    return "未歸類"


def copy_row_style(ws, src_row, dst_row, max_col):
    """複製上一列的儲存格字型、對齊、邊框與數值格式"""
    if src_row < 2:
        return
    for col in range(1, max_col + 1):
        src_cell = ws.cell(src_row, col)
        dst_cell = ws.cell(dst_row, col)
        if src_cell.has_style:
            dst_cell.font = copy(src_cell.font)
            dst_cell.alignment = copy(src_cell.alignment)
            dst_cell.border = copy(src_cell.border)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.number_format = src_cell.number_format


def append_scenarios(excel_path, new_items_json):
    wb = openpyxl.load_workbook(excel_path)
    ws_main = wb["AI應用場景"]
    ws_src = wb["來源索引"]

    # 找出「AI應用場景」最後一筆有效資料列（避免空白列影響）
    last_main_row = 1
    existing = set()
    for r in range(2, ws_main.max_row + 1):
        comp = str(ws_main.cell(r, 2).value or "").strip()  # 第 2 欄 (B): 業者
        scen = str(ws_main.cell(r, 4).value or "").strip()  # 第 4 欄 (D): 場景
        if comp or scen:
            last_main_row = r
        if comp and scen:
            existing.add((comp, scen))

    # 找出「來源索引」左側表格 (A~D 欄) 的最後一筆有效資料列
    last_src_row = 1
    known_companies = set()
    for r in range(2, ws_src.max_row + 1):
        src_comp = str(ws_src.cell(r, 2).value or "").strip()
        if src_comp:
            last_src_row = r
            known_companies.add(src_comp)

    taipei_today = get_taipei_today()
    taipei_today_str = taipei_today.strftime("%Y-%m-%d")

    added_count = 0
    for item in new_items_json:
        sector = str(item.get("sector") or item.get("業別") or "").strip()
        company = str(item.get("company") or item.get("業者") or "").strip()
        raw_holding = str(item.get("holding") or item.get("所屬金控") or "").strip()
        scenario = str(item.get("scenario") or item.get("場景") or "").strip()
        category = str(item.get("category") or item.get("類別") or "").strip()
        raw_date = item.get("update_date") or item.get("更新日期") or item.get("date")

        if not (sector and company and scenario and category):
            continue

        # 1. 判定所屬金控（非屬表列 14 家金控則填入 "未歸類"）
        holding = resolve_holding(company, raw_holding)

        # 2. 取得執行搜尋當天日期（台灣台北時區 Asia/Taipei）
        update_dt = parse_update_date(raw_date, taipei_today)

        # 排除重複之 (業者, 場景)
        if (company, scenario) not in existing:
            next_main_row = last_main_row + 1
            ws_main.cell(next_main_row, 1, sector)
            ws_main.cell(next_main_row, 2, company)
            ws_main.cell(next_main_row, 3, holding)
            ws_main.cell(next_main_row, 4, scenario)
            ws_main.cell(next_main_row, 5, category)
            ws_main.cell(next_main_row, 6, update_dt)

            # 複製上一列樣式並設定格式
            copy_row_style(ws_main, last_main_row, next_main_row, 6)
            ws_main.cell(next_main_row, 3).number_format = "@"
            ws_main.cell(next_main_row, 6).number_format = "yyyy/mm/dd"

            last_main_row = next_main_row
            existing.add((company, scenario))
            added_count += 1

            # 檢查該業者是否已在「來源索引」中，若無則新增一列並帶入 COUNTIFS 公式
            if company not in known_companies:
                next_src_row = last_src_row + 1
                ws_src.cell(next_src_row, 1, sector)
                ws_src.cell(next_src_row, 2, company)
                ws_src.cell(next_src_row, 3, "最新蒐集")
                ws_src.cell(
                    next_src_row,
                    4,
                    f"=COUNTIFS(AI應用場景!A:A,A{next_src_row},AI應用場景!B:B,B{next_src_row})",
                )
                copy_row_style(ws_src, last_src_row, next_src_row, 4)
                last_src_row = next_src_row
                known_companies.add(company)

    # 同步擴展 Excel 表格範圍與資料驗證範圍，確保「分類說明」之結構化公式自動納入新資料
    if "金融業AI場景明細" in ws_main.tables:
        tbl_main = ws_main.tables["金融業AI場景明細"]
        tbl_main.ref = f"A1:F{last_main_row}"
        if tbl_main.autoFilter:
            tbl_main.autoFilter.ref = f"A1:F{last_main_row}"

    for dv in ws_main.data_validations.dataValidation:
        sqref_str = str(dv.sqref)
        for col_letter in ("A", "B", "C", "E"):
            if sqref_str.startswith(f"{col_letter}2:"):
                dv.sqref = f"{col_letter}2:{col_letter}{last_main_row}"

    if "業者來源索引" in ws_src.tables:
        tbl_src = ws_src.tables["業者來源索引"]
        tbl_src.ref = f"A1:D{last_src_row}"
        if tbl_src.autoFilter:
            tbl_src.autoFilter.ref = f"A1:D{last_src_row}"

    wb.save(excel_path)
    print(
        f"成功新增 {added_count} 筆 AI 場景至 {excel_path}（更新日期：{taipei_today_str}，台北時區）"
    )


if __name__ == "__main__":
    # 接收 Agent 傳入的 JSON 字串：
    # python3 update_excel.py '[{"sector":"證券期貨業","company":"凱基證券","holding":"凱基金","scenario":"...","category":"生成式AI"}]'
    items = json.loads(sys.argv[1])
    append_scenarios("金融業AI場景應用.xlsx", items)